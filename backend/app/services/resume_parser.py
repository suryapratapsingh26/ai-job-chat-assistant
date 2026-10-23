"""In-memory PDF/DOCX extraction and deterministic resume profile parsing."""

import re
import unicodedata
import zipfile
from io import BytesIO
from pathlib import Path

from docx import Document
from pypdf import PdfReader

from backend.app.models.resume import ResumeProfile
from backend.app.services.normalization import normalize_skills

MAX_RESUME_BYTES = 10 * 1024 * 1024
MAX_DOCX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_RESUME_PAGES = 100

MIME_TYPES = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}
SECTION_ALIASES = {
    "skills": {"skills", "technical skills", "core competencies", "key skills"},
    "experience": {"experience", "work experience", "professional experience", "employment history", "work history"},
    "education": {"education", "academic background", "certifications", "certificates"},
    "projects": {"projects", "personal projects", "selected projects", "project experience"},
}
TITLE_WORDS = (
    "engineer", "developer", "designer", "manager", "analyst", "consultant", "scientist",
    "architect", "specialist", "lead", "director", "intern", "administrator", "product",
)
EMAIL_RE = re.compile(r"\b[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}\b", re.IGNORECASE)
PHONE_RE = re.compile(r"(?<!\w)(?:\+\s?\d{1,3}[\s().-]?)?(?:\(?\d{2,4}\)?[\s().-]?)?\d{3,4}[\s.-]?\d{3,4}(?!\w)")


class ResumeProcessingError(ValueError):
    def __init__(self, message: str, status_code: int = 422):
        super().__init__(message)
        self.status_code = status_code


def _section_name(line: str):
    candidate = re.sub(r"\s*[:\-–—]\s*$", "", line.strip()).casefold()
    for name, aliases in SECTION_ALIASES.items():
        if candidate in aliases:
            return name, ""
    inline = re.match(r"^\s*([^:]{1,40})\s*:\s*(.+)$", line)
    if inline:
        heading = inline.group(1).strip().casefold()
        for name, aliases in SECTION_ALIASES.items():
            if heading in aliases:
                return name, inline.group(2).strip()
    return None, ""


def _sections(lines: list[str]) -> dict[str, list[str]]:
    found: dict[str, list[str]] = {key: [] for key in SECTION_ALIASES}
    current = None
    for line in lines:
        section, inline_content = _section_name(line)
        if section:
            current = section
            if inline_content:
                found[current].append(inline_content)
            continue
        if current and line.strip():
            found[current].append(line.strip())
    return found


def _clean_lines(text: str) -> list[str]:
    normalized = unicodedata.normalize("NFKC", text).replace("\x00", "")
    lines = []
    for raw_line in normalized.splitlines():
        line = " ".join(raw_line.replace("\t", " ").split()).strip(" •·▪◦-–—\t")
        if line:
            lines.append(line)
    return lines


def _header_lines(lines: list[str]) -> list[str]:
    header = []
    for line in lines:
        if _section_name(line)[0]:
            break
        header.append(line)
        if len(header) == 12:
            break
    return header


def _find_name(lines: list[str]):
    for line in _header_lines(lines)[:8]:
        explicit = re.match(r"(?i)^name\s*:\s*(.+)$", line)
        candidate = explicit.group(1) if explicit else line
        candidate = EMAIL_RE.sub("", candidate)
        candidate = PHONE_RE.sub("", candidate)
        candidate = re.split(r"\s+[|•]\s+|\s{2,}", candidate)[0].strip(" |,;–—-")
        words = candidate.split()
        if not 2 <= len(words) <= 5:
            continue
        if any(word.casefold() in {"resume", "curriculum", "vitae", "profile", "portfolio"} for word in words):
            continue
        if any(word.casefold().rstrip(".,") in TITLE_WORDS for word in words):
            continue
        if all(re.fullmatch(r"[^\W\d_][^\d_]*[.]?", word, re.UNICODE) for word in words):
            return candidate
    return None


def _find_phone(text: str):
    for match in PHONE_RE.finditer(text):
        value = " ".join(match.group().split())
        digits = re.sub(r"\D", "", value)
        if re.fullmatch(r"(?:19|20)\d{2}[\s–-]+(?:19|20)\d{2}", value):
            continue
        if 8 <= len(digits) <= 15:
            return value
    return None


def _find_designation(lines: list[str]):
    header = _header_lines(lines)
    for line in header:
        explicit = re.match(r"(?i)^(?:designation|current role|current position|role|title)\s*:\s*(.+)$", line)
        if explicit:
            return explicit.group(1).strip()[:120] or None
    for line in header:
        lowered = line.casefold()
        if any(word in lowered for word in TITLE_WORDS) and len(line) <= 120 and "@" not in line:
            return line[:120]
    return None


def parse_profile(text: str) -> ResumeProfile:
    """Extract details present in the resume; absent values remain empty."""
    lines = _clean_lines(text)
    if not lines:
        raise ResumeProcessingError("No readable text was found in this resume.")

    sections = _sections(lines)
    skill_candidates = []
    for line in sections["skills"]:
        skill_candidates.extend(re.split(r"[,;|•·]", re.sub(r"^[^:]{1,30}:\s*", "", line)))

    education = [line[:500] for line in sections["education"][:20]]
    projects = [line[:500] for line in sections["projects"][:20]]
    experience = "\n".join(sections["experience"])[:6000] or None
    email_match = EMAIL_RE.search(text)
    return ResumeProfile(
        name=_find_name(lines),
        email=email_match.group(0) if email_match else None,
        phone=_find_phone(text),
        designation=_find_designation(lines),
        skills=normalize_skills(skill_candidates)[:100],
        experience=experience,
        education=education,
        projects=projects,
    )


def _extract_pdf(data: bytes) -> str:
    if not data.lstrip().startswith(b"%PDF-"):
        raise ResumeProcessingError("This PDF file is invalid or damaged.")
    try:
        reader = PdfReader(BytesIO(data), strict=False)
        if reader.is_encrypted:
            raise ResumeProcessingError("This PDF is password-protected. Upload an unlocked copy.")
        if len(reader.pages) > MAX_RESUME_PAGES:
            raise ResumeProcessingError(f"This PDF has more than {MAX_RESUME_PAGES} pages. Upload a shorter resume.", 413)
        extracted = "\n".join(page.extract_text() or "" for page in reader.pages)
    except ResumeProcessingError:
        raise
    except Exception as exc:
        raise ResumeProcessingError("Could not read this PDF. Try exporting it again or upload a DOCX file.") from exc
    return extracted


def _extract_docx(data: bytes) -> str:
    try:
        with zipfile.ZipFile(BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > 2000 or sum(item.file_size for item in entries) > MAX_DOCX_UNCOMPRESSED_BYTES:
                raise ResumeProcessingError("This DOCX expands beyond the supported size limit.", 413)
            names = {item.filename for item in entries}
            if "[Content_Types].xml" not in names or "word/document.xml" not in names:
                raise ResumeProcessingError("This DOCX file is invalid or damaged.")
        document = Document(BytesIO(data))
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        table_text = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    except ResumeProcessingError:
        raise
    except Exception as exc:
        raise ResumeProcessingError("Could not read this DOCX file. Try saving it again in Microsoft Word.") from exc
    return "\n".join([*paragraphs, *table_text])


def parse_resume_document(filename: str, content_type: str, data: bytes) -> ResumeProfile:
    if len(data) > MAX_RESUME_BYTES:
        raise ResumeProcessingError("Resume file exceeds the 10 MB size limit.", 413)
    extension = Path(filename or "").suffix.casefold()
    if extension not in MIME_TYPES:
        raise ResumeProcessingError("Unsupported file type. Upload a PDF or DOCX resume.", 415)
    mime = (content_type or "").split(";", 1)[0].strip().casefold()
    if mime and mime not in {MIME_TYPES[extension], "application/octet-stream"}:
        raise ResumeProcessingError("The file type does not match its extension. Upload a PDF or DOCX resume.", 415)

    if extension == ".pdf":
        text = _extract_pdf(data)
    else:
        if not data.startswith(b"PK"):
            raise ResumeProcessingError("This DOCX file is invalid or damaged.")
        text = _extract_docx(data)

    if not text.strip():
        raise ResumeProcessingError("No extractable text was found. Scanned or image-only resumes are not supported yet.")
    return parse_profile(text)
